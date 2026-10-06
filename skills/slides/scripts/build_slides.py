#!/usr/bin/env python3
"""Build editable 16:9 slides from a JSON deck spec (slides skill).

Usage:
    python build_slides.py deck.json -o deck.pptx

The deck spec is described in references/deck-spec.md. Every value comes from
the spec; anything left as "[__]" is reported as a missing slot.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# --- Style tokens (keep in sync with references/style.md) -------------------

NAVY = RGBColor(0x1F, 0x3A, 0x6E)
BLUE = RGBColor(0x3B, 0x5F, 0xC4)
TEAL = RGBColor(0x2A, 0x8A, 0x9A)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
ORANGE = RGBColor(0xC0, 0x58, 0x2E)
PAGE = RGBColor(0xF6, 0xF7, 0xF9)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xD9, 0xDE, 0xE6)
ROW_ALT = RGBColor(0xF0, 0xF3, 0xF7)
TEXT = RGBColor(0x1C, 0x24, 0x33)
MUTED = RGBColor(0x4D, 0x56, 0x66)
FOOTER = RGBColor(0x6A, 0x72, 0x82)
PANEL_BLUE = RGBColor(0xE8, 0xEE, 0xFB)
PANEL_GREEN = RGBColor(0xE5, 0xF3, 0xF0)
AXIS = RGBColor(0x9A, 0xA4, 0xB5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ACCENTS = [NAVY, TEAL, BLUE, GREEN]
TONES = {
    "good": GREEN, "attention": ORANGE, "neutral": NAVY, "proposed": BLUE,
    "navy": NAVY, "blue": BLUE, "teal": TEAL, "green": GREEN, "orange": ORANGE,
}
PANELS = {"blue": PANEL_BLUE, "green": PANEL_GREEN}
GOOD_WORDS = ("on track", "ongoing", "done", "complete", "approved", "go-ahead", "passed")
ATTENTION_WORDS = ("pending", "at risk", "blocked", "delayed", "failed", "open")

FONT = "Open Sans"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.67
HEADER_H = 1.17
ACCENT_H = 0.06
TOP = 1.5
BOTTOM = 6.4
CONTENT_W = SLIDE_W - 2 * MARGIN
FOOTER_Y = SLIDE_H - 0.44 - 0.28
MISSING = "[__]"


# --- Low-level helpers ------------------------------------------------------

def tone(name, default=NAVY):
    return TONES.get(str(name or "").lower(), default)


def auto_tone(text):
    t = str(text).lower()
    if any(w in t for w in GOOD_WORDS):
        return GREEN
    if any(w in t for w in ATTENTION_WORDS):
        return ORANGE
    return None


def _font(run, size, color, bold):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rpr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def _runs(p, text, size, color, bold=False):
    """Add text to a paragraph; **double asterisks** mark bold spans."""
    parts = re.split(r"\*\*(.+?)\*\*", str(text))
    for i, part in enumerate(parts):
        if part:
            run = p.add_run()
            run.text = part
            _font(run, size, color, bold or i % 2 == 1)


def textbox(slide, x, y, w, h, lines, anchor="t"):
    """lines: list of dicts {text, size, color, bold, align, before, after}."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right"):
        setattr(tf, f"margin_{side}", Inches(0.04))
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf._txBody.bodyPr.set("anchor", anchor)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(line.get("before", 0))
        p.space_after = Pt(line.get("after", 2))
        _runs(p, line.get("text", MISSING), line.get("size", 13), line.get("color", TEXT), line.get("bold", False))
    return box


def rect(slide, x, y, w, h, fill=None, line=None, line_pt=1, shape=MSO_SHAPE.RECTANGLE, radius=None, dash=False):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_pt)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    shp.shadow.inherit = False
    return shp


def card(slide, x, y, w, h, accent=None, side="left", thick=0.055):
    rect(slide, x, y, w, h, CARD, BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    if accent is not None:
        if side == "left":
            rect(slide, x, y, thick, h, accent)
        else:
            rect(slide, x, y, w, thick, accent)


def panel(slide, x, y, w, h, color="blue", border=None):
    rect(slide, x, y, w, h, PANELS.get(color, PANEL_BLUE), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    if border is not None:
        rect(slide, x, y, 0.07, h, border)


def line(slide, x1, y1, x2, y2, color, pt=1.5, dash=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(pt)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def heading_text(heading, text, head_size=14, body_size=13, head_color=NAVY):
    return [
        {"text": heading, "size": head_size, "color": head_color, "bold": True, "after": 6},
        {"text": text, "size": body_size},
    ]


def _cell_value(value):
    if isinstance(value, dict):
        return str(value.get("text", MISSING)), value.get("tone"), value.get("bold")
    return str(value), None, None


def table(slide, x, y, w, columns, rows, header_fills=None, widths=None, status_col=None,
          numeric=False, row_h=0.42, size=12):
    """rows: list of lists, or dicts {cells, key, indent}. Cells may be {text, tone, bold}."""
    n_rows = max(len(rows), 1)
    shape = slide.shapes.add_table(n_rows + 1, len(columns), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * (n_rows + 1)))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    widths = widths or [1 / len(columns)] * len(columns)
    for i, frac in enumerate(widths):
        tbl.columns[i].width = Inches(w * frac)
    for r in range(n_rows + 1):
        tbl.rows[r].height = Inches(row_h)

    def put(cell, text, color, bold, fill, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        tf = cell.text_frame
        tf.word_wrap = True
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = align
        _runs(p, text, size, color, bold)

    for c, label in enumerate(columns):
        fill = header_fills[c % len(header_fills)] if header_fills else NAVY
        align = PP_ALIGN.RIGHT if numeric and c > 0 else PP_ALIGN.LEFT
        put(tbl.cell(0, c), label, WHITE, True, fill, align)

    for r, row in enumerate(rows or [[MISSING] * len(columns)], start=1):
        meta = row if isinstance(row, dict) else {"cells": row}
        fill = ROW_ALT if r % 2 == 0 else CARD
        for c in range(len(columns)):
            raw = meta["cells"][c] if c < len(meta["cells"]) else MISSING
            text, tname, bold = _cell_value(raw)
            color = tone(tname) if tname else TEXT
            if c == status_col and not tname:
                color = auto_tone(text) or TEXT
                bold = True if bold is None else bold
            if numeric and c == 1 and meta.get("key"):
                color, bold = BLUE, True
            if c == 0 and meta.get("indent"):
                text = "    " + text
            align = PP_ALIGN.RIGHT if numeric and c > 0 else PP_ALIGN.LEFT
            put(tbl.cell(r, c), text, color, bool(bold), fill, align)
    return row_h * (n_rows + 1)


# --- Frame ------------------------------------------------------------------

def frame(slide, spec, page, footer):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, PAGE)
    rect(slide, 0, 0, SLIDE_W, HEADER_H, NAVY)
    rect(slide, 0, HEADER_H, SLIDE_W, ACCENT_H, BLUE)
    pill = spec.get("pill")
    textbox(slide, MARGIN, 0.25, 9.6 if pill else 12.0, 0.68,
            [{"text": spec.get("title", MISSING), "size": 30, "color": WHITE, "bold": True}], anchor="ctr")
    if pill:
        label = pill.get("text", MISSING)
        pw = max(1.6, 0.11 * len(label) + 0.7)
        px = SLIDE_W - MARGIN - pw
        rect(slide, px, 0.37, pw, 0.44, tone(pill.get("tone"), BLUE), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        textbox(slide, px, 0.37, pw, 0.44,
                [{"text": label, "size": 13, "color": WHITE, "bold": True, "align": PP_ALIGN.CENTER}], anchor="ctr")
    foot = " · ".join(t for t in (footer, spec.get("source_note")) if t)
    if foot:
        textbox(slide, MARGIN, FOOTER_Y, 9.5, 0.32, [{"text": foot, "size": 12, "color": FOOTER}])
    textbox(slide, SLIDE_W - MARGIN - 1.2, FOOTER_Y, 1.2, 0.32,
            [{"text": str(page), "size": 12, "color": FOOTER, "align": PP_ALIGN.RIGHT}])
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


# --- Layouts ----------------------------------------------------------------

def table_kpi(slide, s):
    cols = s.get("columns", [MISSING])
    rows = s.get("rows", [])
    status_col = next((i for i, c in enumerate(cols) if "status" in str(c).lower()), None)
    h = table(slide, MARGIN, TOP, CONTENT_W, cols, rows, header_fills=ACCENTS, widths=s.get("widths"),
              status_col=status_col, row_h=0.5)
    kpis = s.get("kpis", [])
    if not kpis:
        return
    y = min(TOP + h + 0.3, BOTTOM - 1.3)
    gap = 0.18
    w = (CONTENT_W - gap * (len(kpis) - 1)) / len(kpis)
    for i, k in enumerate(kpis):
        x = MARGIN + i * (w + gap)
        card(slide, x, y, w, 1.3, ACCENTS[i % 4], side="top", thick=0.07)
        textbox(slide, x + 0.16, y + 0.2, w - 0.3, 1.0, [
            {"text": k.get("value", MISSING), "size": 26, "color": NAVY, "bold": True, "after": 4},
            {"text": k.get("label", MISSING), "size": 12, "color": MUTED},
        ])


def scorecard(slide, s):
    lw, gap = 3.9, 0.22
    rx = MARGIN + lw + gap
    rw = SLIDE_W - MARGIN - rx
    has_panels = bool(s.get("panels"))
    hero = s.get("hero", {})
    htone = tone(hero.get("tone"), GREEN)
    card(slide, MARGIN, TOP, lw, 1.95, htone, thick=0.07)
    textbox(slide, MARGIN + 0.2, TOP + 0.1, lw - 0.3, 1.75, [
        {"text": hero.get("label", MISSING), "size": 14, "color": NAVY, "bold": True, "after": 2},
        {"text": hero.get("value", MISSING), "size": 60, "color": htone, "bold": True, "after": 2},
        {"text": hero.get("sub_label", ""), "size": 12, "color": MUTED},
    ], anchor="ctr")
    chips = s.get("chips", [])
    if chips:
        cg = 0.1
        cw = (lw - cg * (len(chips) - 1)) / len(chips)
        for i, c in enumerate(chips):
            x = MARGIN + i * (cw + cg)
            ct = tone(c.get("tone"), NAVY)
            card(slide, x, TOP + 2.1, cw, 1.0, ct, side="top", thick=0.05)
            textbox(slide, x + 0.04, TOP + 2.18, cw - 0.08, 0.86, [
                {"text": c.get("value", MISSING), "size": 26, "color": ct, "bold": True, "align": PP_ALIGN.CENTER},
                {"text": c.get("label", MISSING), "size": 11, "color": MUTED, "align": PP_ALIGN.CENTER},
            ], anchor="ctr")
    if s.get("detail"):
        textbox(slide, MARGIN, TOP + 3.2, lw, 0.6, [{"text": s["detail"], "size": 12, "color": MUTED}])
    img_h = 3.3 if has_panels else BOTTOM - TOP - 0.45
    img = s.get("image")
    if img and Path(img).exists():
        from PIL import Image
        with Image.open(img) as im:
            iw, ih = im.size
        scale = min(rw / iw, img_h / ih)
        pw, ph = iw * scale, ih * scale
        slide.shapes.add_picture(img, Inches(rx + (rw - pw) / 2), Inches(TOP + (img_h - ph) / 2), Inches(pw), Inches(ph))
    else:
        rect(slide, rx, TOP, rw, img_h, CARD, BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
        textbox(slide, rx, TOP + img_h / 2 - 0.3, rw, 0.6,
                [{"text": "[image]", "size": 16, "color": FOOTER, "align": PP_ALIGN.CENTER}], anchor="ctr")
    if s.get("caption"):
        textbox(slide, rx, TOP + img_h + 0.05, rw, 0.32, [{"text": s["caption"], "size": 12, "color": FOOTER}])
    if has_panels:
        pans = s["panels"]
        pw = (CONTENT_W - gap * (len(pans) - 1)) / len(pans)
        for i, p in enumerate(pans):
            x = MARGIN + i * (pw + gap)
            panel(slide, x, 5.15, pw, 1.2, p.get("tone", "blue" if i % 2 == 0 else "green"))
            textbox(slide, x + 0.18, 5.25, pw - 0.36, 1.0, heading_text(p.get("heading", MISSING), p.get("text", MISSING)))


def data_table_insights(slide, s):
    lw = CONTENT_W * 0.6
    gap = 0.22
    rx = MARGIN + lw + gap
    rw = SLIDE_W - MARGIN - rx
    h = table(slide, MARGIN, TOP, lw, s.get("columns", [MISSING]), s.get("rows", []), numeric=True, row_h=0.4)
    if s.get("summary"):
        textbox(slide, MARGIN, min(TOP + h + 0.15, BOTTOM - 0.45), lw, 0.45,
                [{"text": s["summary"], "size": 12, "color": MUTED}])
    cards = s.get("cards", [])
    if cards:
        cg = 0.14
        ch = (BOTTOM - TOP - cg * (len(cards) - 1)) / len(cards)
        for i, c in enumerate(cards):
            y = TOP + i * (ch + cg)
            card(slide, rx, y, rw, ch, tone(c.get("tone"), ACCENTS[i % 4]), thick=0.07)
            textbox(slide, rx + 0.22, y + 0.12, rw - 0.36, ch - 0.2, heading_text(c.get("heading", MISSING), c.get("text", MISSING)), anchor="ctr")


def step_flow(slide, s):
    y = TOP
    if s.get("intro"):
        textbox(slide, MARGIN, y, CONTENT_W, 0.5, [{"text": s["intro"], "size": 16}])
        y += 0.6
    steps = s.get("steps", [])
    aw = 0.3
    sw = (CONTENT_W - aw * (len(steps) - 1)) / max(len(steps), 1)
    sh = 1.85
    for i, st in enumerate(steps):
        x = MARGIN + i * (sw + aw)
        col = ACCENTS[i % 4]
        card(slide, x, y, sw, sh, col, side="top", thick=0.07)
        textbox(slide, x + 0.14, y + 0.2, sw - 0.26, sh - 0.3, [
            {"text": st.get("label", MISSING), "size": 13, "color": col, "bold": True, "after": 6},
            {"text": st.get("text", MISSING), "size": 13},
        ])
        if i < len(steps) - 1:
            rect(slide, x + sw + 0.06, y + sh / 2 - 0.1, 0.18, 0.2, FOOTER, shape=MSO_SHAPE.RIGHT_ARROW)
    y += sh + 0.25
    pans = s.get("panels", [])
    if pans:
        gap = 0.2
        pw = (CONTENT_W - gap * (len(pans) - 1)) / len(pans)
        ph = 1.45
        for i, p in enumerate(pans):
            x = MARGIN + i * (pw + gap)
            panel(slide, x, y, pw, ph, p.get("tone", "green" if i % 2 == 0 else "blue"))
            textbox(slide, x + 0.18, y + 0.12, pw - 0.36, ph - 0.2, heading_text(p.get("heading", MISSING), p.get("text", MISSING)))
        y += ph + 0.2
    if s.get("closing"):
        textbox(slide, MARGIN, min(y, BOTTOM - 0.4), CONTENT_W, 0.4,
                [{"text": s["closing"], "size": 13, "color": NAVY, "bold": True}])


def numbered_cards(slide, s):
    textbox(slide, MARGIN, TOP, CONTENT_W, 0.85, [{"text": s.get("headline", MISSING), "size": 17}])
    cards = s.get("cards", [])
    gap = 0.18
    cw = (CONTENT_W - gap * (len(cards) - 1)) / max(len(cards), 1)
    y, ch = TOP + 1.0, 2.55
    for i, c in enumerate(cards):
        x = MARGIN + i * (cw + gap)
        card(slide, x, y, cw, ch, NAVY, side="top", thick=0.07)
        textbox(slide, x + 0.18, y + 0.2, cw - 0.32, ch - 0.3, [
            {"text": str(c.get("number", i + 1)), "size": 13, "color": BLUE, "bold": True, "after": 4},
            {"text": c.get("heading", MISSING), "size": 15, "color": NAVY, "bold": True, "after": 6},
            {"text": c.get("text", MISSING), "size": 13},
        ])
    note = s.get("note")
    if note:
        ny = y + ch + 0.25
        panel(slide, MARGIN, ny, CONTENT_W, 0.75, "blue")
        lead = f"**{note['lead_in']}** " if note.get("lead_in") else ""
        textbox(slide, MARGIN + 0.2, ny + 0.08, CONTENT_W - 0.4, 0.6,
                [{"text": lead + note.get("text", MISSING), "size": 13}], anchor="ctr")


def _d(v):
    return date.fromisoformat(v)


def milestone_timeline(slide, s):
    start, end = _d(s["start"]), _d(s["end"])
    span = max((end - start).days, 1)
    x0, x1 = MARGIN + 0.35, SLIDE_W - MARGIN - 0.35

    def xpos(v):
        return x0 + (x1 - x0) * (_d(v) - start).days / span

    axis_y = 3.85
    for ph in s.get("phases", []):
        a, b = xpos(ph["start"]), xpos(ph["end"])
        rect(slide, a, TOP + 0.05, max(b - a, 1.4), 0.36, PANEL_BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        textbox(slide, a + 0.08, TOP + 0.05, max(b - a, 1.4) - 0.1, 0.36,
                [{"text": ph.get("label", MISSING), "size": 13, "color": NAVY, "bold": True}], anchor="ctr")
    line(slide, MARGIN, axis_y, SLIDE_W - MARGIN, axis_y, AXIS, pt=2.5)
    today = s.get("today")
    if today:
        tx = xpos(today["date"])
        line(slide, tx, TOP + 0.5, tx, axis_y + 0.25, ORANGE, pt=2, dash=True)
        textbox(slide, tx - 1.25, TOP + 0.48, 1.2, 0.3,
                [{"text": today.get("label", MISSING), "size": 12, "color": ORANGE, "bold": True, "align": PP_ALIGN.RIGHT}])
    kinds = {"event": NAVY, "approval": GREEN, "target": BLUE, "golive": GREEN}
    ms = sorted(s.get("milestones", []), key=lambda m: m["date"])
    lw = 1.65
    for i, m in enumerate(ms):
        x = xpos(m["date"])
        kind = m.get("kind", "event")
        col = kinds.get(kind, NAVY)
        if kind == "golive":
            rect(slide, x - 0.13, axis_y - 0.13, 0.26, 0.26, col, shape=MSO_SHAPE.DIAMOND)
        else:
            rect(slide, x - 0.09, axis_y - 0.09, 0.18, 0.18, col, shape=MSO_SHAPE.OVAL)
        lx = min(max(x - lw / 2, MARGIN), SLIDE_W - MARGIN - lw)
        above = i % 2 == 0
        ly = axis_y - 1.25 if above else axis_y + 0.25
        tcol = col if kind != "event" else TEXT
        textbox(slide, lx, ly, lw, 1.0, [
            {"text": m.get("label", MISSING), "size": 12, "color": tcol, "bold": True, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": m.get("event", MISSING), "size": 12, "color": tcol, "align": PP_ALIGN.CENTER},
        ], anchor="b" if above else "t")
    y = axis_y + 1.4
    if s.get("pattern"):
        card(slide, MARGIN, y, CONTENT_W, 0.55, NAVY, thick=0.07)
        textbox(slide, MARGIN + 0.25, y, CONTENT_W - 0.4, 0.55,
                [{"text": s["pattern"], "size": 14, "color": NAVY, "bold": True}], anchor="ctr")
        y += 0.65
    if s.get("caveat"):
        textbox(slide, MARGIN, y, CONTENT_W, 0.35, [{"text": s["caveat"], "size": 12, "color": FOOTER}])


def lane_timeline(slide, s):
    start, end = _d(s["start"]), _d(s["end"])
    span = max((end - start).days, 1)
    label_w = 1.9
    x0, x1 = MARGIN + label_w, SLIDE_W - MARGIN

    def xpos(v):
        return x0 + (x1 - x0) * (_d(v) - start).days / span

    lanes = s.get("lanes", [])
    lane_h = min(0.8, (BOTTOM - TOP - 1.0) / max(len(lanes), 1))
    top = TOP + 0.4
    bottom = top + lane_h * len(lanes)
    for li in range(1, len(lanes), 2):
        rect(slide, MARGIN, top + li * lane_h, CONTENT_W, lane_h, ROW_ALT)
    for p in s.get("periods", []):
        px = xpos(p["date"])
        line(slide, px, top, px, bottom, BORDER, pt=0.75)
        textbox(slide, px + 0.03, TOP, 1.4, 0.32, [{"text": p.get("label", MISSING), "size": 11, "color": MUTED}])
    for li, lane in enumerate(lanes):
        y = top + li * lane_h
        textbox(slide, MARGIN, y, label_w - 0.1, lane_h,
                [{"text": lane.get("label", MISSING), "size": 13, "color": NAVY, "bold": True}], anchor="ctr")
        bh = lane_h * 0.5
        by = y + (lane_h - bh) / 2
        for bi, bar in enumerate(lane.get("bars", [])):
            a, b = xpos(bar["start"]), xpos(bar["end"])
            w = max(b - a, 0.08)
            col = tone(bar.get("tone"), ACCENTS[(li + bi) % 4])
            style = bar.get("style", "done")
            if style == "wait":
                shp = rect(slide, a, by, w, bh, None, col, line_pt=0.75)
                shp.fill.patterned()
                shp.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
                shp.fill.fore_color.rgb = col
                shp.fill.back_color.rgb = WHITE
            elif style == "planned":
                rect(slide, a, by, w, bh, WHITE, col, line_pt=1.5, dash=True)
            else:
                rect(slide, a, by, w, bh, col)
            label = bar.get("label")
            if label:
                inside = style == "done" and w > 0.12 * len(label)
                textbox(slide, a + 0.05 if inside else a + w + 0.05, by, max(w - 0.1, 2.4), bh,
                        [{"text": label, "size": 11, "color": WHITE if inside else TEXT, "bold": inside}], anchor="ctr")
        for m in lane.get("markers", []):
            mx = xpos(m["date"])
            current = m.get("kind") == "current"
            rect(slide, mx - 0.13, by + bh / 2 - 0.13, 0.26, 0.26, GREEN if current else WHITE,
                 None if current else MUTED, line_pt=1.5, shape=MSO_SHAPE.DIAMOND)
            if m.get("label"):
                textbox(slide, mx - 0.8, by + bh / 2 + 0.14, 1.6, 0.3,
                        [{"text": m["label"], "size": 11, "color": GREEN if current else MUTED, "align": PP_ALIGN.CENTER}])
            if m.get("note"):
                textbox(slide, mx - 2.2, by - 0.32, 2.1, 0.3,
                        [{"text": m["note"], "size": 11, "color": ORANGE, "bold": True, "align": PP_ALIGN.RIGHT}])
    today = s.get("today")
    if today:
        tx = xpos(today["date"])
        line(slide, tx, top - 0.1, tx, bottom + 0.05, ORANGE, pt=2, dash=True)
        textbox(slide, tx - 0.6, bottom + 0.05, 1.2, 0.3,
                [{"text": today.get("label", MISSING), "size": 11, "color": ORANGE, "bold": True, "align": PP_ALIGN.CENTER}])
    legend = s.get("legend")
    if legend:
        lx, ly = MARGIN + label_w, bottom + 0.45
        for style, text in legend.items():
            if style == "wait":
                shp = rect(slide, lx, ly + 0.06, 0.4, 0.2, None, NAVY, line_pt=0.75)
                shp.fill.patterned()
                shp.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
                shp.fill.fore_color.rgb = NAVY
                shp.fill.back_color.rgb = WHITE
            elif style == "planned":
                rect(slide, lx, ly + 0.06, 0.4, 0.2, WHITE, NAVY, dash=True)
            else:
                rect(slide, lx, ly + 0.06, 0.4, 0.2, NAVY)
            textbox(slide, lx + 0.48, ly, 2.4, 0.32, [{"text": text, "size": 11, "color": MUTED}])
            lx += 3.0


def status_learnings(slide, s):
    lw = CONTENT_W * 0.4
    gap = 0.35
    rx = MARGIN + lw + gap
    rw = SLIDE_W - MARGIN - rx
    textbox(slide, MARGIN, TOP, lw, 0.4, [{"text": s.get("left_heading", MISSING), "size": 15, "color": NAVY, "bold": True}])
    cols = s.get("columns", [MISSING, MISSING])
    h = table(slide, MARGIN, TOP + 0.45, lw, cols, s.get("rows", []), widths=[0.52, 0.48], status_col=1, row_h=0.4)
    path = s.get("path")
    if path:
        py = TOP + 0.45 + h + 0.2
        panel(slide, MARGIN, py, lw, 1.0, "green", border=GREEN)
        lead = f"**{path['lead_in']}** " if path.get("lead_in") else ""
        textbox(slide, MARGIN + 0.2, py + 0.08, lw - 0.3, 0.85, [{"text": lead + path.get("text", MISSING), "size": 12}], anchor="ctr")
    textbox(slide, rx, TOP, rw, 0.4, [{"text": s.get("right_heading", MISSING), "size": 15, "color": NAVY, "bold": True}])
    items = s.get("learnings", [])
    cg = 0.1
    ch = min(0.78, (BOTTOM - 0.5 - TOP - 0.45 - cg * (len(items) - 1)) / max(len(items), 1))
    for i, it in enumerate(items):
        y = TOP + 0.45 + i * (ch + cg)
        rect(slide, rx, y, rw, ch, CARD, BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        textbox(slide, rx + 0.18, y + 0.04, rw - 0.3, ch - 0.08, [
            {"text": f"{i + 1} · {it.get('title', MISSING)}", "size": 13, "color": BLUE, "bold": True, "after": 1},
            {"text": it.get("text", MISSING), "size": 12},
        ], anchor="ctr")
    if s.get("closing"):
        textbox(slide, MARGIN, BOTTOM - 0.35, CONTENT_W, 0.4,
                [{"text": s["closing"], "size": 13, "color": NAVY, "bold": True}])


LAYOUTS = {
    "table-kpi": table_kpi,
    "scorecard": scorecard,
    "data-table-insights": data_table_insights,
    "step-flow": step_flow,
    "numbered-cards": numbered_cards,
    "milestone-timeline": milestone_timeline,
    "lane-timeline": lane_timeline,
    "status-learnings": status_learnings,
}
ALIASES = {"status-table": "table-kpi", "metrics-table": "data-table-insights", "proposal-flow": "step-flow"}


# --- Deck -------------------------------------------------------------------

def missing_slots(obj, path="slides"):
    found = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            found += missing_slots(v, f"{path}.{k}")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            found += missing_slots(v, f"{path}[{i}]")
    elif isinstance(obj, str) and MISSING in obj:
        found.append(path)
    return found


def build(deck: dict) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    for page, spec in enumerate(deck.get("slides", []), start=deck.get("first_page", 1)):
        name = ALIASES.get(spec.get("layout"), spec.get("layout"))
        if name not in LAYOUTS:
            raise SystemExit(f"Unknown layout '{spec.get('layout')}'. Use one of: {', '.join(LAYOUTS)}")
        slide = prs.slides.add_slide(blank)
        frame(slide, spec, page, deck.get("footer"))
        LAYOUTS[name](slide, spec)
    return prs


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="deck spec JSON file")
    ap.add_argument("-o", "--output", help="output .pptx (default: next to the spec)")
    args = ap.parse_args()
    spec_path = Path(args.spec)
    deck = json.loads(spec_path.read_text())
    out = Path(args.output) if args.output else spec_path.with_suffix(".pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    build(deck).save(out)
    print(f"Saved {out}")
    gaps = missing_slots(deck.get("slides", []))
    if gaps:
        print("Missing slots ([__]):")
        for g in gaps:
            print(f"  - {g}")


if __name__ == "__main__":
    main()
